"""Measure frozen synthetic proposal/review, cold CLI, and resume budgets."""
from __future__ import annotations

import contextlib
import hashlib
import inspect
import io
import json
import math
import platform
import socket
import statistics
import subprocess
import sys
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ats_agent.capture import capture_cached, fetch_many  # noqa: E402
from ats_agent.formatting import audit_file  # noqa: E402
from ats_agent.ingestion import load  # noqa: E402
from ats_agent.providers import DeterministicRewriteProvider, RewriteCache  # noqa: E402
from ats_agent.workflow import apply_manifest, build_proposal  # noqa: E402
from ats_agent.review import build_approval_manifest, write_review_bundle  # noqa: E402
from ats_agent.orchestrator import tailor  # noqa: E402
from ats_agent.doctor import _doctor  # noqa: E402

REPETITIONS = 20
FIXTURES = ROOT / "benchmarks/perf-fixtures/workflow-v2"


def _freeze_fixtures() -> dict:
    """Freeze varied synthetic inputs once, then reject any hash drift."""
    manifest = FIXTURES / "manifest.json"
    if not manifest.exists():
        FIXTURES.mkdir(parents=True, exist_ok=True)
        if any(FIXTURES.iterdir()):
            raise ValueError("incomplete performance fixtures; inspect before recreating")
        matrix = json.loads((ROOT / "benchmarks/v3/public-development.jsonl").read_text().splitlines()[0])
        families = list(matrix["families"].values())
        aliases = matrix["aliases"]

        def sentence(index: int) -> str:
            family = families[index % len(families)]
            term = family["terms"][(index // len(families)) % len(family["terms"])]
            domain = family["domains"][index % len(family["domains"])]
            return (f"Helped configure {aliases.get(term, term)} for {domain}; documented input assumptions, "
                    "reviewed exceptions with the project team, and checked the resulting outputs "
                    "against the agreed requirements before sharing the analysis.")

        document = Document()
        style = document.styles["Normal"]
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(0)
        document.add_heading("Synthetic Candidate", level=1)
        document.add_heading("Experience", level=2)
        for index in range(28):
            if index == 14:
                document.add_page_break()
                document.add_heading("Projects", level=2)
            document.add_paragraph(sentence(index))
        document.save(FIXTURES / "resume.docx")
        for index in range(5):
            (FIXTURES / f"evidence-{index}.md").write_text(
                "\n".join(sentence(28 + index * 18 + line) for line in range(18)), encoding="utf-8")
        for index in range(20):
            family = families[index % len(families)]
            lines = [family["roles"][index]]
            for line in range(20):
                term = family["terms"][(index + line) % len(family["terms"])]
                domain = family["domains"][(index + line) % len(family["domains"])]
                lines.append(f"Experience with {term} is preferred for {domain}. "
                             "The analyst reviews source material, checks assumptions with colleagues, "
                             "documents the reasoning, and explains the completed work clearly.")
            (FIXTURES / f"job-{index:02d}.md").write_text("\n".join(lines), encoding="utf-8")
        hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(FIXTURES.iterdir())}
        manifest.write_text(json.dumps({"schema_version": 1, "files": hashes}, indent=2), encoding="utf-8")
    expected = json.loads(manifest.read_text())["files"]
    for name, digest in expected.items():
        if hashlib.sha256((FIXTURES / name).read_bytes()).hexdigest() != digest:
            raise ValueError(f"frozen performance fixture changed: {name}")
    return expected


def _summary(samples: list[float], *, digits: int = 4) -> dict:
    return {"p50": round(statistics.median(samples), digits),
            "p95": round(sorted(samples)[math.ceil(len(samples) * 0.95) - 1], digits)}


class _CountingProvider(DeterministicRewriteProvider):
    """Count deterministic calls without introducing a live or paid provider."""

    def __init__(self) -> None:
        self.calls = 0
        self.latencies: list[float] = []

    def generate(self, context):  # noqa: ANN001
        self.calls += 1
        started = time.perf_counter()
        try:
            return super().generate(context)
        finally:
            self.latencies.append(time.perf_counter() - started)


def _outcome_counts(result: dict) -> tuple[dict[str, int], int]:
    """Return visible status counts and the number of blocked roles."""

    counts: dict[str, int] = {}
    for outcome in result.get("outcomes", []):
        status = str(outcome.get("status", "unknown"))
        counts[status] = counts.get(status, 0) + 1
    return counts, sum(count for status, count in counts.items() if status.startswith("blocked"))


def main() -> int:
    hashes = _freeze_fixtures()
    resume = FIXTURES / "resume.docx"
    evidence = [FIXTURES / f"evidence-{index}.md" for index in range(5)]
    samples: dict[str, list[float]] = {name: [] for name in (
        "warm_prepare", "time_to_first_proposal", "ten_roles", "twenty_roles",
        "analysis", "review", "cold_prepare", "capture", "capture_cache_hit", "apply", "validation",
        "resume_twenty", "sequential_distinct_host_capture", "parallel_distinct_host_capture")}
    changes = 0
    with tempfile.TemporaryDirectory(prefix="ats-performance-") as temporary:
        out = Path(temporary)
        supports_cache = "candidate_cache" in inspect.signature(build_proposal).parameters
        for repetition in range(REPETITIONS):
            cache: dict = {}
            start = time.perf_counter()
            analysis_seconds = review_seconds = 0.0
            for index in range(20):
                before = time.perf_counter()
                proposal = build_proposal(
                    resume,
                    FIXTURES / f"job-{index:02d}.md",
                    evidence_paths=evidence,
                    candidate_id="synthetic-perf",
                    **({"candidate_cache": cache} if supports_cache else {}),
                )
                analysis_seconds += time.perf_counter() - before
                if proposal.get("status") != "draft" or not proposal.get("requirements"):
                    raise ValueError("performance input failed to produce a substantive draft")
                changes += sum(bool(c.get("supported")) for c in proposal["changes"])
                before = time.perf_counter()
                write_review_bundle(proposal, out / "reviews" / str(index))
                review_seconds += time.perf_counter() - before
                if index in {0, 9, 19}:
                    samples[{0: "warm_prepare", 9: "ten_roles", 19: "twenty_roles"}[index]].append(
                        time.perf_counter() - start
                    )
                if index == 0:
                    samples["time_to_first_proposal"].append(analysis_seconds)
            samples["analysis"].append(analysis_seconds)
            samples["review"].append(review_seconds)
            before = time.perf_counter()
            command = [sys.executable, "-m", "ats_agent.cli", "prepare", str(resume),
                       str(FIXTURES / "job-00.md"), "--candidate-id", "synthetic-perf",
                       "--out", str(out / f"cold-{repetition}")]
            for path in evidence:
                command.extend(["--evidence", str(path)])
            subprocess.run(command, cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)
            samples["cold_prepare"].append(time.perf_counter() - before)

        provider = _CountingProvider()
        provider_cache = RewriteCache()
        for _ in range(REPETITIONS):
            duplicate = build_proposal(
                resume,
                FIXTURES / "job-00.md",
                evidence_paths=evidence,
                candidate_id="synthetic-perf",
                provider=provider,
                candidate_cache={},
                rewrite_cache=provider_cache,
            )
            if duplicate.get("status") != "draft":
                raise ValueError("provider-deduplication input failed to produce a draft")
        provider_calls = {"observed": provider.calls, **provider_cache.stats()}
        provider_latency = _summary(provider.latencies, digits=6)
        if provider_calls["cache_hits"] == 0:
            raise ValueError("identical provider contexts were not deduplicated")

        for repetition in range(REPETITIONS):
            proposal = build_proposal(
                resume,
                FIXTURES / "job-00.md",
                evidence_paths=evidence,
                candidate_id="synthetic-perf",
            )
            supported = next(
                (
                    change
                    for change in proposal.get("changes", [])
                    if change.get("supported") and change.get("operation") == "replace_span"
                ),
                None,
            )
            if supported is None:
                raise ValueError("performance fixture has no strict-preserve rewrite")
            stage = out / f"delivery-{repetition}"
            stage.mkdir()
            proposal_path = stage / "proposal.json"
            proposal_path.write_text(json.dumps(proposal, indent=2) + "\n", encoding="utf-8")
            manifest = build_approval_manifest(
                proposal,
                proposal_filename=proposal_path.name,
                selections=[(
                    supported["id"],
                    supported.get("default_variant") or supported["variants"][0]["id"],
                )],
                output_document="tailored.docx",
                document_mode="strict-preserve",
            )
            approval_path = stage / "approval.json"
            approval_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
            before = time.perf_counter()
            applied = apply_manifest(approval_path)
            samples["apply"].append(time.perf_counter() - before)
            output = Path(str(applied["output"]))
            before = time.perf_counter()
            parsed = load(output)
            formatting = audit_file(str(output))
            samples["validation"].append(time.perf_counter() - before)
            if not parsed["body_text"] or formatting.get("status") != "audited":
                raise ValueError("synthetic output validation failed")

        source = json.dumps([{"id": f"role-{index}", "job_url": f"https://perf.example/role-{index}"}
                             for index in range(20)])

        def capture(url, destination, *, source_type):
            time.sleep(0.05)  # Controlled latency for capture/resume comparisons; never live network.
            index = int(url.rsplit("-", 1)[1])
            body = (FIXTURES / f"job-{index:02d}.md").read_bytes()
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_bytes(body)
            return {"url": url, "path": str(destination), "sha256": hashlib.sha256(body).hexdigest(),
                    "captured_at": datetime.now(timezone.utc).isoformat(), "source_type": source_type,
                    "method": "benchmark_mock", "extraction_status": "captured"}

        with patch("ats_agent.capture.socket.getaddrinfo", return_value=[
            (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443)),
        ]), patch("ats_agent.capture.capture_url", side_effect=capture) as fetch, \
             patch("ats_agent.orchestrator.capture_url", new=fetch, create=True), \
             patch("ats_agent.orchestrator.build_proposal", wraps=build_proposal) as propose, \
             contextlib.redirect_stderr(io.StringIO()):
            capture_start = fetch.call_count
            capture_cache_hits = 0
            for repetition in range(REPETITIONS):
                before = time.perf_counter()
                captured = capture_cached(
                    f"https://perf.example/capture-{repetition}",
                    out / f"capture-cache-{repetition}",
                    source_type="job_posting",
                    destination=out / f"capture-{repetition}.txt",
                )
                samples["capture"].append(time.perf_counter() - before)
                if captured.get("cache") != "miss":
                    raise ValueError("fresh capture unexpectedly reused a cache entry")
                before = time.perf_counter()
                cached = capture_cached(
                    f"https://perf.example/capture-{repetition}",
                    out / f"capture-cache-{repetition}",
                    source_type="job_posting",
                    destination=out / f"capture-hit-{repetition}.txt",
                )
                samples["capture_cache_hit"].append(time.perf_counter() - before)
                if cached.get("cache") != "hit":
                    raise ValueError("fresh capture did not produce a reusable cache entry")
                capture_cache_hits += 1
            capture_calls = fetch.call_count - capture_start
            concurrent_capture_start = fetch.call_count
            concurrency_ratios: list[float] = []
            for repetition in range(REPETITIONS):
                urls = [
                    (f"https://perf-{index}.example/role-{index}",
                     out / f"sequential-{repetition}-{index}.txt")
                    for index in range(4)
                ]
                before = time.perf_counter()
                sequential = fetch_many(urls, source_type="job_posting", workers=1, use_cache=False)
                sequential_elapsed = time.perf_counter() - before
                samples["sequential_distinct_host_capture"].append(sequential_elapsed)
                parallel_urls = [
                    (url, out / f"parallel-{repetition}-{index}.txt")
                    for index, (url, _destination) in enumerate(urls)
                ]
                before = time.perf_counter()
                parallel = fetch_many(parallel_urls, source_type="job_posting", workers=4, use_cache=False)
                parallel_elapsed = time.perf_counter() - before
                samples["parallel_distinct_host_capture"].append(parallel_elapsed)
                if not all(isinstance(item, dict) for item in [*sequential, *parallel]):
                    raise ValueError("distinct-host mock capture unexpectedly blocked a role")
                concurrency_ratios.append(parallel_elapsed / sequential_elapsed)
            concurrent_capture_calls = fetch.call_count - concurrent_capture_start
            before_initial_tailor = fetch.call_count
            initial_result = tailor(resume, source, evidence_paths=evidence, candidate_id="synthetic-perf",
                                    run_dir=out / "resume", writer=lambda _: None)
            initial_tailor_fetches = fetch.call_count - before_initial_tailor
            initial_fetches, initial_proposals = fetch.call_count, propose.call_count
            initial_outcomes, initial_blocked = _outcome_counts(initial_result)
            for _ in range(REPETITIONS):
                before = time.perf_counter()
                result = tailor(resume, source, evidence_paths=evidence, candidate_id="synthetic-perf",
                                run_dir=out / "resume", writer=lambda _: None)
                samples["resume_twenty"].append(time.perf_counter() - before)
                if len(result["outcomes"]) != 20:
                    raise ValueError("resume dropped a role")
            resume_fetches = fetch.call_count - initial_fetches
            resume_proposals = propose.call_count - initial_proposals
            resume_outcomes, resume_blocked = _outcome_counts(result)
    timing = {name: _summary(values) for name, values in samples.items()}
    concurrency_ratio = _summary(concurrency_ratios)
    budgets = {
        "warm_prepare": 2.0,
        "cold_prepare": 3.0,
        "ten_roles": 15.0,
        "twenty_roles": 30.0,
        "resume_twenty": 2.0,
    }
    passed = (all(timing[name]["p95"] <= limit for name, limit in budgets.items())
              and capture_calls == REPETITIONS and resume_fetches == 0
              and resume_proposals == 0 and changes > 0
              and concurrency_ratio["p95"] <= 0.5
              and initial_blocked == 0 and resume_blocked == 0)
    print(json.dumps({
        "schema_version": 5, "measurement": "synthetic_offline_with_review_and_mocked_capture",
        "repetitions": REPETITIONS, "roles": 20, "timing_seconds": timing,
        "budgets_seconds": budgets, "supported_changes_observed": changes,
        "call_counts": {
            "mock_capture": capture_calls,
            "mock_capture_cache_hits": capture_cache_hits,
            "mock_distinct_host_capture": concurrent_capture_calls,
            "initial_tailor_capture": initial_tailor_fetches,
            "resume_capture": resume_fetches,
            "resume_proposal": resume_proposals,
            "provider": provider_calls,
        },
        "outcomes": {
            "initial": initial_outcomes,
            "initial_blocked": initial_blocked,
            "resume": resume_outcomes,
            "resume_blocked": resume_blocked,
        },
        "provider_latency_seconds": provider_latency,
        "distinct_host_concurrency": {
            "workers": 4,
            "paired_parallel_to_sequential_ratio": concurrency_ratio,
            "target_ratio": 0.5,
        },
        "capture_note": "mocked; no live-site or paid-provider latency measured",
        "mock_capture_latency_seconds": 0.05,
        "layout_note": "explicit page break; rendered pagination unverified",
        "fixture_sha256": hashes, "runtime": _doctor().get("runtime"),
        "environment": {"python": sys.version.split()[0], "platform": platform.platform()},
        "status": "passed" if passed else "failed",
    }, indent=2))
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(main())
